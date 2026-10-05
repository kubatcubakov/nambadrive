"""Opt-in real Tika / OpenSearch contracts. Never connect to production indexes."""

import io
import os
import uuid
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pydantic import SecretStr

from app.core.config import Settings
from app.search.clients import OpenSearch, Tika
from app.search.service import search
from tests.test_documents import clean as clean_fixture
from tests.test_documents import scene as scene_fixture

clean, scene = clean_fixture, scene_fixture


def test_real_tika_office_text_and_no_ocr():
    url = os.environ.get("NAMBADRIVE_TEST_TIKA_URL")
    if not url:
        pytest.skip("Requires isolated real Tika")
    tika = Tika(Settings(tika_url=url))
    word = Document()
    word.add_paragraph("Namba confidential contract")
    docx = io.BytesIO()
    word.save(docx)
    workbook = Workbook()
    workbook.active.append(["Namba", "contract"])
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    for content, mime in [
        (docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        (io.BytesIO(b"Namba confidential contract"), "text/plain"),
        (io.BytesIO(b"Namba,contract\nA,1"), "text/csv"),
        (io.BytesIO(b'{"Namba": "contract"}'), "application/json"),
        (io.BytesIO(b"<Namba>contract</Namba>"), "application/xml"),
    ]:
        result = tika.extract(content, mime)
        assert "contract" in result, mime
    # Image-only PDF must not silently invoke OCR.
    pdf = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(pdf, format="PDF")
    assert not tika.extract(pdf, "application/pdf").strip()


async def test_real_opensearch_acl_filter(db, scene, clean):
    url = os.environ.get("NAMBADRIVE_TEST_OPENSEARCH_URL")
    if not url:
        pytest.skip("Requires isolated disposable OpenSearch")
    settings = Settings(
        opensearch_url=url,
        search_index="nd-ci-" + uuid.uuid4().hex,
        opensearch_username=os.environ.get("NAMBADRIVE_TEST_SEARCH_USER", "isolated-test"),
        opensearch_password=SecretStr(
            os.environ.get("NAMBADRIVE_TEST_SEARCH_PASSWORD", "isolated-test")
        ),
        opensearch_ca_file=os.environ.get("NAMBADRIVE_TEST_SEARCH_CA"),
    )
    engine = OpenSearch(settings)
    row, version, _ = clean
    try:
        engine.initialize()
        engine.put(row.id, version.id, row.name, "Namba confidential contract")
        engine.request("POST", "/_refresh")
        candidates = engine.candidates("contract")
        assert len(candidates) == 1 and "contract" in candidates[0].snippet
        with patch("app.search.service.write_audit_event", new=AsyncMock()):
            assert await search(db, scene[1], engine, "contract", 20, {}) == []
            assert len(await search(db, scene[0], engine, "contract", 20, {})) == 1
            version.is_current = False
            await db.commit()
            assert await search(db, scene[0], engine, "contract", 20, {}) == []
        engine.delete(row.id)
        engine.request("POST", "/_refresh")
        assert engine.candidates("contract") == []
    finally:
        engine.request("DELETE", "")
