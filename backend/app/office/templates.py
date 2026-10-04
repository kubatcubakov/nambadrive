from __future__ import annotations

import io


def blank(extension: str) -> io.BytesIO:
    output = io.BytesIO()
    if extension == "docx":
        from docx import Document

        Document().save(output)
    elif extension == "xlsx":
        from openpyxl import Workbook  # type: ignore[import-untyped]

        workbook = Workbook()
        workbook.save(output)
        workbook.close()
    elif extension == "pptx":
        from pptx import Presentation

        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(output)
    else:
        raise ValueError("Unsupported Office template")
    output.seek(0)
    return output
